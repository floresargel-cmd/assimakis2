from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

content = '''
# Startup and connections
Prevents another instance with the same executable filename from starting.
Shows a splash screen and opens maximized.
Connects to crane 2, the old oven, and the new oven.
Opens the basket database and EPICS database; production startup fails if either connection cannot open.
Uses connection settings compiled into the application.
Creates missing application log, mission, and export tables.
Trims each application log table to its latest 10,000 records at startup.
Loads stored basket positions and checks for unfinished missions.
# Main display
Displays oven positions, buffers, return positions, conveyors, destacker positions, and crane carriage.
Shows basket numbers, contents, destination classifications, export assignments, and locks.
Highlights mission source and destination positions.
Displays crane coordinates, targets, current mission, next mission, and PLC step.
Provides mouse-wheel zoom, middle-button pan, and Zoom Extents.
Shows detailed basket/profile information in tooltips.
Updates the window title with the current time.
# Operating modes
Starts in Missions mode.
Provides Missions, Exports, and Modify selection modes.
Provides separate export lists for HCB/SB, anodizing, packing destacker 1, and packing destacker 2.
Selecting the automatic-controls page switches layout selection back to Missions mode.
Automatic scheduling options start disabled.
# Manual crane missions
Clicking an eligible occupied position selects the source.
Clicking an eligible empty position attempts the move immediately.
Clicking the selected source or another occupied position cancels source selection.
Escape cancels a pending selection.
Rejects moves when the crane is unavailable, communications report an error, the source is empty, or the destination is occupied.
Rejects locked, obstructed, unsupported, or already-reserved positions.
Persists accepted missions and sends their endpoints to the PLC.
# Stack and obstruction rules
Prevents moving a basket with another basket immediately above it.
Prevents using an upper level when the level immediately below is empty.
Checks configured obstructions from adjacent oven columns.
Reserves entire position columns against conflicting missions.
Excludes old-oven columns 660 and 680 from normal manual source/destination selection.
Represents three usable levels in each buffer and three return-stack levels.
# Mission execution and tracking
Maintains one next mission and one active mission.
Promotes the next mission when the PLC reports starting.
Moves the recorded basket from source to carriage when the PLC reports pickup.
Moves it from carriage to destination when the PLC reports completion.
Acknowledges mission stages to the PLC.
Clears completed mission records.
Logs mismatches such as an occupied destination or missing carriage basket.
Checks mission status every second.
# Automatic scheduling
Supports independent enable switches for four export categories, north imports, south imports, and EPICS queue population.
Evaluates missions every 2.5 seconds and immediately when a mode is enabled.
Considers the first basket in each enabled export queue.
Uses this fixed priority: HCB/SB, anodizing, packing 1, packing 2, north import, then south import.
Sends all export categories to the north conveyor handoff.
Imports north/south baskets into the return stack.
Checks route occupancy and downstream packing workload.
Does not automatically try a lower-priority candidate when the selected candidate fails final mission validation.
# Automatic uncovering
Detects baskets blocking a queued export basket.
Attempts to relocate a blocker into an available north or south buffer position.
Searches buffer levels from lower to higher.
Moves a blocker already in one buffer only to the other buffer.
Leaves the exported basket queued while relocating blockers.
# EPICS basket information
Reads profile, quantity, length, temper, sales order, order line, finish code, standard practice, customer, next department, and production stop time.
Loads contents for basket/rack numbers 1–130 with nonzero quantities.
Converts order length to millimetres.
Classifies contents using department, finish, customer, and standard-practice text.
Classifies non-DESTACK contents as None.
Routes matching finishing/customer/practice cases to HCB; remaining DESTACK cases become HCA.
Uses the designated top contents to classify the basket.
Can populate export queues from the highest occupied basket in each eligible oven/buffer column.
Uses matching sales order and line to continue selected destacker workloads.
# Export list management
Lets operators assign baskets to one of four export queues.
Manual reassignment removes the basket from its previous export queues.
Appends new entries at the end of the queue.
Lets operators move entries up/down or remove them.
Shows queue counts and an empty-list message.
Refreshes lists every five seconds.
Manual packing assignment requires maximum profile length ≤7660 mm.
Removes completed north-handoff baskets from crane-2 export queues.
Adds qualifying packing baskets to the corresponding crane-3 queue.
# Conveyor routing
Writes a destination code based on export membership.
Supports None, Anodizing, Destacker For SB, To Press, In from Destacker, Anodizing Overlength, and Packing Destacker.
Contains routing thresholds of 7821 mm for HCB/anodizing and 7660 mm for packing.
Logs destination writes to a local text file.
Does not automatically clear an existing destination when the basket has no export-queue membership.
# Returning baskets to crane 1
Fills return positions from bottom to top.
Requests return automatically after a completed delivery fills all three levels.
Provides a toolbar action to request return earlier.
Sends the bottom basket number to the old-oven PLC.
Does not itself clear all three basket locations in that request function.
# Basket and position correction
Ctrl-click opens position information where supported.
Shift+left-click reloads position data.
Allows basket-number and position-lock editing.
Uses basket number 0 to represent no basket.
Shows basket contents and movement history.
Blocks basket-number changes for positions participating in stored missions.
Attempts to prevent declaring a basket already located elsewhere.
Directs most PLC-controlled position edits to the conveyor interface.
Conveyor editing writes basket number and destination directly to the PLC.
# Crane HMI controls
Provides Auto, Manual, Active Stop, and Reset Alarms.
Provides X/Y jogging and Z up/down controls.
Provides basket-pin and safety-pin lock/unlock.
Displays position, target, velocity, frequency, current, proximity, and on-target feedback.
Provides axis trend plots.
Provides Continue, Abort, End, Clear Next, and Clear Active Database Mission.
Provides encoder resets and an interlock-bypass command.
F1 sends the Active Stop command.
Basic HMI provides alarm reset and alarm/warning displays.
# Conveyor HMI controls
Displays eight conveyor positions and their proximity switches.
Provides conveyor Auto, Manual, and alarm reset.
Provides individual forward/backward jogging.
Provides paired conveyor movement commands.
Provides conveyor rotation, lift up/down, and side movement.
Graphical movement buttons send a command on press and zero on release.
Displays active conveyor commands and alarms.
# Destacker HMI controls
Displays health feedback for eight pincers.
Provides individual pincer enable/disable commands.
Provides a Spacers on Belts control.
Displays destacker alarms.
# Recovery access and logging
Warns about a stored active mission at startup and offers to clear it.
Attempts to restore a stored next mission.
Attempts to return recorded basket locations to the source after abort.
Password-protects the PLC communication dialog and unlocking selected advanced mission buttons.
Initially disables Abort, End, and Clear Active Database controls whenever the full HMI opens.
Opens the full HMI and ordinary position editor without a password.
Records general, verbose, GUI, position, and configured PLC-message history.
Shows recent logs and larger history dialogs.
Consolidates consecutive identical GUI messages by updating their timestamp.
# Current defects affecting behavior
Settings password enforcement is broken: the result is discarded by a comma expression.
The additional password check inside Clear Active Database Mission has the same defect, although its advanced-button unlock remains separate.
A bConveyorState/BConveyorState mismatch prevents the intended handoff-readiness update through the shown path.
Handoff routing can evaluate the destination's previous contents length before loading the arriving basket.
An immediate export-marker helper incorrectly maps packing destacker 2.
Automatic uncovering can run for an individually disabled export category when another automatic option is enabled.
Maintenance mode is forced true in the conveyor update handler.
Stored next missions can reject themselves during restart validation.
Abort recovery may require another status-processing cycle before clearing.
Duplicate-location detection misses baskets recorded in multiple locations.
The crane drawing's Z value uses a different PLC index from the HMI.
Profile top and bottom ordering is inconsistent between helpers.
The current production EPICS classifier never returns anodizing, despite having an automatic anodizing branch.
Failure to open the local destination log exits the application.
Database transfers and queue updates use multiple calls without an explicit transaction in this application layer.
'''

root=Path(__file__).resolve().parent.parent
doc=Document()
for level in doc.part.numbering_part.element.iter(qn('w:lvl')):
    fmt=level.find(qn('w:numFmt'))
    if fmt is not None and fmt.get(qn('w:val')) == 'bullet':
        level.find(qn('w:lvlText')).set(qn('w:val'), '\u2022')
        rpr=level.find(qn('w:rPr'))
        if rpr is None:
            rpr=OxmlElement('w:rPr'); level.append(rpr)
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts=OxmlElement('w:rFonts'); rpr.append(fonts)
        fonts.set(qn('w:ascii'),'Calibri'); fonts.set(qn('w:hAnsi'),'Calibri')
for element in list(doc.styles.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
    st=doc.styles[name]
    st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].paragraph_format.space_after=Pt(6)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Subtitle'].font.italic=False
for element in list(doc.styles['Subtitle'].element.iter(qn('w:spacing'))):
    if element.getparent().tag == qn('w:rPr'):
        element.getparent().remove(element)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)
h=doc.styles['Heading 1']; h.font.size=Pt(14); h.font.bold=True
h.paragraph_format.space_before=Pt(12); h.paragraph_format.space_after=Pt(6)
h.paragraph_format.keep_with_next=True
b=doc.styles['List Bullet']; b.font.size=Pt(11)
b.paragraph_format.space_after=Pt(3)
b.paragraph_format.line_spacing=1.03
b.paragraph_format.keep_together=True
doc.add_paragraph('Apex Basket Crane 2', 'Title')
doc.add_paragraph('Current application behavior', 'Subtitle')
doc.add_paragraph('Code analysis dated September 25 2026')
doc.add_paragraph('The application controls basket crane 2 workflows, maintains basket locations, and coordinates exports, imports, conveyors, and destacker operations. This document lists its current behavior, followed by code defects that affect those workflows.')
doc.add_paragraph('Scope: the checked-in apexBasketCrane2 solution and its related source. This is static analysis; the application was not run and no equipment or databases were connected. PLC commands describe requests made by the application, not verified physical machine actions.')
num=0; bullets=0
for line in content.strip().splitlines():
    if line.startswith('# '):
        num+=1
        doc.add_paragraph(f'{num}  {line[2:]}','Heading 1')
    else:
        doc.add_paragraph(line,'List Bullet'); bullets+=1
doc.add_paragraph('Inactive production workflows','Heading 1')
doc.add_paragraph('Emulation, legacy crane functionality, and the unused generic connection-settings dialog are not active production workflows in the checked-in configuration.')
doc.add_paragraph('Detailed source reference','Heading 1')
doc.add_paragraph('The companion APP_BEHAVIOR_INVENTORY.md contains 182 detailed entries, source references, PLC command mappings, and 320 configured alarm and status labels.')
p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Page '); r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.core_properties.title='Apex Basket Crane 2 Current Application Behavior'
doc.core_properties.subject='Current code based application behavior'
doc.core_properties.author=''
out=root/'Apex_Basket_Crane_2_Current_Behavior.docx'
doc.save(out)
print(f'{out}\nSections: {num}; behavior bullets: {bullets}')
